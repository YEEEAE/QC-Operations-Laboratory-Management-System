#!/usr/bin/env python3
"""Author SOP-FRS-014 review draft from the SOP-FRS-011 family template.

Preserves the template styles, header/footer, numbering, theme and section
properties; replaces body content; embeds live screenshots and Stitch visual
proposals. Review-draft only: no approval, effectivity or in-application
signature. Proposed UX improvements are NOT IMPLEMENTED.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/approved/SOP-FRS-011-D0.1-Workspace-Tools.docx"
OUT = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/review-drafts/SOP-FRS-014-D0.1-User-Experience-All-Pages.docx"
ATTACH = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/review-drafts/SOP-FRS-014-attachments"
LIVE = ATTACH / "live"
STITCH = ATTACH / "stitch"

HEADER = "REVIEW DRAFT — not approved | yazeed (prepared) | 09 October 2026 | Effective date not assigned"
FOOTER = "SOP-FRS-014 | D0.1 | REVIEW DRAFT | Page "


def clear_body(doc: Document):
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        body.remove(sect)
    for child in list(body):
        body.remove(child)
    return sect


def set_header_footer(doc: Document):
    for section in doc.sections:
        for p in section.header.paragraphs:
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = HEADER
            else:
                p.add_run(HEADER)
        for p in section.footer.paragraphs:
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = FOOTER
            else:
                p.add_run(FOOTER)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if "USING WORKSPACE TOOLS" in p.text:
                            for r in p.runs:
                                r.text = r.text.replace("USING WORKSPACE TOOLS", "USER EXPERIENCE — ALL PAGES")
                        if "SOP-FRS-011" in p.text:
                            for r in p.runs:
                                r.text = r.text.replace("SOP-FRS-011", "SOP-FRS-014")


def para(doc, text, style="Normal"):
    return doc.add_paragraph(text, style=style)


def figure(doc, image: Path, caption: str):
    if not image.exists():
        raise FileNotFoundError(image)
    from PIL import Image

    with Image.open(image) as im:
        w_px, h_px = im.size
    max_w, max_h = 16.0, 16.0
    h_at_max_w = max_w * h_px / w_px
    if h_at_max_w <= max_h:
        doc.add_picture(str(image), width=Cm(max_w))
    else:
        doc.add_picture(str(image), height=Cm(max_h))
    para(doc, caption)


# route, section kicker, h1, purpose, primary action, controls/notes
PAGE_UX = [
    ("/dashboard", "Overview", "QC OPERATIONAL COMMAND CENTER", "Dashboard",
     "No create action; start-of-day surface.",
     "Needs-your-attention queue, decision work queues, eleven operational action counts each with source/scope and a drill-down link, the six-stage quarantine flow, a receiving-per-day series, and a Data coverage panel. A server snapshot time and Refresh sit under the heading."),
    ("/tasks", "Work", "Work queue", "Tasks",
     "Create task.",
     "GET filter bar (State, Due, Assignee, Show, Sort by, Direction) with Apply filters; bounded pagination with a true total (page 1 of 1 at zero); an empty state that offers the create action."),
    ("/work", "Work", "Your queue", "My work today",
     "No create action; grouped personal queue.",
     "Four groups (Blocked, Overdue, Due today, Assigned to you) each with its own empty state and a published limit; groups may overlap and a record appears once."),
    ("/quality/findings", "Quality", "Findings", "Findings",
     "Create finding.",
     "State filter plus Apply filter; a filtered empty state ('No findings match this view') distinct from a genuinely empty register."),
    ("/quality/ncr", "Quality", "Quality register", "NCR records",
     "No create action while governed decisions are open.",
     "A blocker banner names PD-15, PD-16 and PD-17 with their owner and keeps actions unavailable while existing records stay readable; State filter."),
    ("/quality/rca", "Quality", "Quality register", "Root-cause analyses",
     "No create action while governed decisions are open.",
     "Purpose line 'Root cause analyses linked to NCRs.'; blocker banner for PD-17; State filter."),
    ("/quality/capa", "Quality", "Quality register", "CAPA records",
     "No create action while governed decisions are open.",
     "Blocker banner for PD-18 plus a plain-language statement of the approved P-04 Supervisor-close exception; State filter."),
    ("/reject-reports", "Quality", "Quality", "Reject reports",
     "Create report.",
     "Purpose line 'Issue slips record approval confirmations. Daily rejects need no approval.'; links to the analytics and daily views."),
    ("/quarantine", "Quarantine", "Operations", "Quarantine",
     "Open the receiving records behind each stage.",
     "Purpose line 'Receiving, inspections and release status.'; six ordered stages, each opening the receiving register with the filter that reproduces its own count."),
    ("/quarantine/receiving", "Quarantine", "Operations", "Receiving items",
     "Create receiving item.",
     "Purpose line 'Inspection result and release state are shown separately.'; the densest filter bar in the product (search, item, lot, supplier, purchase order, workflow state, inspection status, quarantine status, inspection result, release state, received/expiry ranges, ownership, received-on) with Apply filters; an unbounded register that declares it shows every item in scope."),
    ("/quarantine/inspections", "Quarantine", "Controlled execution", "Inspection reports",
     "Open the receiving register.",
     "Purpose line 'Report status and scientific result are shown separately.'; the register separates report workflow state from the scientific result."),
    ("/quarantine/admin", "Quarantine", "Inspection templates", "Quarantine administration",
     "Create template.",
     "Purpose line 'Manage inspection templates. Approved content requires a new revision to change.'; when the catalogue is unavailable the page keeps the register readable and shows a reload hint before creating."),
    ("/laboratory/tests", "Laboratory", "CONTROLLED LABORATORY", "Laboratory tests",
     "Create laboratory test.",
     "Purpose line names the two transcription-only report drafts; GET filter bar (Workflow state, Ownership, Sort by, Direction) with Apply filters and bounded pagination; a clear no-records state with the create action."),
    ("/assets/equipment", "Assets", "Equipment", "Equipment",
     "Add equipment.",
     "Purpose line 'Stable equipment identity and its current operational state.'; Operational state filter; an unbounded register declaring it shows every record in scope; empty state with the create action."),
    ("/assets/calibrations", "Assets", "Calibrations", "Calibration records",
     "Add calibration record.",
     "Purpose line 'Historical calibration records linked to equipment. Dates are source-provided; no interval is inferred.'; State filter; unbounded register; empty state."),
    ("/assets/maintenance", "Assets", "Maintenance", "Maintenance",
     "Add maintenance record.",
     "Purpose line 'Preserve work history and show equipment impact explicitly. Completion does not prove calibration.'; State filter; unbounded register; empty state."),
    ("/approvals", "Governance", "Approvals", "My approvals",
     "No create action; assigned decisions only.",
     "Purpose line 'Assigned approval work is listed with its current decision status. Permission to review does not by itself make a decision available.'; a signature banner explains that some decisions now require an account-bound signature and warns not to enter a password for a blocked decision; the empty state explicitly says the empty list is not evidence that every decision is available."),
    ("/change-requests", "Governance", "Change requests", "Change requests",
     "New change request.",
     "Purpose line 'Review proposed changes. Approval does not automatically update the record.'; the create flow is contextual for document-version changes."),
    ("/documents", "Governance", "Documents", "Controlled documents",
     "Create document.",
     "Purpose line 'Browse WI, SOP, and other controlled documents available to your account.'; a Type filter and a My-review-queue toggle; an empty state that explains only permitted documents are shown."),
    ("/governance/releases/[releaseId]", "Governance", "Releases", "Release evidence",
     "No editable controls; evidence is read-only.",
     "The release-evidence view is read-only: server-derived evidence, identity and risk are shown without editable checkboxes or JSON."),
    ("/reports", "Insights", "Report list", "Reports",
     "Open report.",
     "The report list names each report, its scope and what it contains; opening a report binds its export to the displayed dataset."),
    ("/reports/[reportCode]", "Insights", "Report", "Report viewer",
     "Download CSV/XLSX or print.",
     "Filter bar, a declared row count and provenance, and request-local export that returns a stale conflict if the dataset changed; print output is labelled informational and unapproved."),
    ("/ai-advisory", "Insights", "AI advisory", "AI advisory",
     "Send only when policy and consent are present.",
     "Purpose line 'Suggestions and analysis only — not an approval authority.'; external processing is disabled by default and the send control stays disabled without an approved policy and per-request consent."),
    ("/admin/users", "Administration", "Members", "Users",
     "Create user.",
     "Search accounts, an Account state filter and Apply filters; a table of login identity, display name, state, roles and scopes; one authorized member account in the current environment."),
    ("/admin/roles", "Administration", "Roles", "Roles",
     "No create action; read-only catalog.",
     "Purpose line 'Read-only role catalog.'; a role table linking to role detail; grants are managed elsewhere."),
    ("/admin/permissions", "Administration", "Permissions", "Permissions",
     "No create action; read-only catalog.",
     "Purpose line 'Read-only catalog. Manage grants through roles.'; the permission catalog is long and intended for lookup."),
    ("/system/health", "System", "(no kicker)", "System health",
     "No create action; owner-only.",
     "Purpose line 'Live system checks.'; separates connectivity from required-workflow readiness and from QC acceptance, shows a scoped identity ratio with a check time, and gives recovery guidance for disabled AI processing. Refused with 404 to any non-owner account."),
    ("/system/control-center", "System", "Owner console", "System control center",
     "Create account; open workspace/roles.",
     "Purpose line 'Manage system access and operations.'; read-only overviews plus the canonical account, role and scope administration actions. Owner-only."),
    ("/system/backups", "System", "Backups", "Backups",
     "No create action; restore is separate.",
     "Purpose line 'Backup records. Restore verification is shown separately.'; publishes the catalog read time and states that an empty catalog, an unavailable provider and unverified recovery are separate states; RPO/RTO show 'Not approved' with the open decision."),
    ("/notifications", "Workspace tools", "(no kicker)", "Notifications",
     "No create action; recipient-scoped.",
     "Purpose line 'Only notifications addressed to your authenticated account are shown.'; an empty state is stated as no notifications sent to this account."),
    ("/search", "Workspace tools", "GLOBAL SEARCH", "Find authorized records",
     "Submit a search query.",
     "Permission- and scope-bounded search; a wordmark kicker and a single search input."),
    ("/account", "Workspace tools", "Your account", "Account",
     "Save password; sign out.",
     "Purpose line 'Manage your password and session.'; changing the password affects active sessions and copy is field-safe."),
    ("/help", "Workspace tools", "Support", "Operating guides",
     "Open a guide or a workflow link.",
     "Purpose line 'Daily workflows and help with blocked actions.'; four guidance tables and deep links into the product, including start-of-day by role."),
    ("/audit", "Audit", "Audit", "Audit history",
     "No create action; read-only history.",
     "Purpose line 'View the audit history available to your account. Use filters to find a record or action.'; a filter bar and a bounded, ordered event table."),
    ("/login", "Entry", "(guest shell)", "Sign in",
     "Sign in.",
     "The only unauthenticated surface; uses an independent animated background and is excluded from the authenticated shell."),
]

IMPROVEMENTS = [
    ("UX-014-01", "Page-anatomy consistency",
     "Pages differ in how they open: some have a section kicker, some none; kicker wording ranges from 'Operations' and 'CONTROLLED LABORATORY' to 'Quality register' and 'Findings'; several pages have no one-line purpose under the heading.",
     "Introduce one page-header primitive that always renders a section kicker from an approved vocabulary, a single H1, an optional one-line purpose and an optional primary action in a fixed order.",
     "Every page becomes predictable to scan and to learn; orientation cost drops.",
     "High", "Design system; delivery pages"),
    ("UX-014-02", "Loading-state continuity",
     "Registers first announce 'Updating results…' and paint a loading shell before the settled register arrives, so the user briefly reads an in-progress state and the layout shifts when content lands.",
     "Render the settled server state on first paint, or a skeleton whose geometry matches the final table, and announce the live region once per navigation rather than per filter.",
     "Removes double-reading and layout shift on every register visit.",
     "High", "Register surfaces; shared DataTable/FilterBar"),
    ("UX-014-03", "Empty vs filtered-empty vocabulary",
     "Empty registers use several different strings ('No matches', 'No records in this view', 'No findings match this view') that do not consistently distinguish a genuinely empty register from a filtered-to-empty one.",
     "Standardize through the shared EmptyTableState tokens: one calm genuinely-empty message with the create action, and one filtered-empty message with a Clear-filters action.",
     "Users immediately know whether to create a record or clear a filter.",
     "Medium", "Shared UI primitives"),
    ("UX-014-04", "Filter-bar progressive disclosure",
     "The receiving register exposes more than a dozen controls at once; other registers show four to six. The primary filters are not visually separated from advanced ones.",
     "Keep the two or three primary filters visible, move the rest behind a More filters disclosure, show applied-filter chips with Clear all, and allow saved presets per user.",
     "Faster triage on dense registers without hiding power.",
     "High", "Register surfaces (Receiving first)"),
    ("UX-014-05", "Bounded registers and true totals",
     "Several registers (receiving, assets, documents, quality) declare 'Showing every … in your authorized scope' with no page control and no total, so findability degrades as records grow.",
     "Push the scope predicate into the query, use the shared pagination with a published true total, and keep the declared-scope wording only where unbounded is deliberate.",
     "Predictable performance and findability at scale.",
     "High", "Register surfaces; companion of IMP-013-04"),
    ("UX-014-06", "Unified policy-blocker card",
     "Policy blockers on NCR, RCA and CAPA render as multi-sentence prose banners with different wording and no common structure.",
     "Use one blocker card primitive showing the decision ID, the owner, a one-line consequence and a link to the governing decision, everywhere an action is unavailable for policy reasons.",
     "The reason an action is missing becomes consistent and findable.",
     "Medium", "Quality pages; companion of IMP-013-05"),
    ("UX-014-07", "Primary-action label consistency",
     "The primary action is labelled differently in the header and the empty state ('Create receiving item' vs 'Create a receiving item'; 'Add equipment' vs 'Create equipment').",
     "Define one canonical label per action and reuse the same string in the header, the empty state and any dialog title.",
     "One action, one name; removes hesitation about whether the two are the same.",
     "Low", "Shared UI copy"),
    ("UX-014-08", "Sticky record header on detail pages",
     "Long record detail pages lose their state, blocking reason and primary action once the reader scrolls into history and evidence.",
     "Add a compact sticky record header carrying the reference, the state and the primary next action, plus an in-page section index for long records.",
     "The next action stays reachable without scrolling back.",
     "Medium", "Record detail surfaces"),
    ("UX-014-09", "Classified feedback and refusal region",
     "Refusals, provider-unavailable states and validation errors appear in several different regions and wordings across registers and detail pages.",
     "Use one classified feedback region that preserves entered values, links errors to fields, distinguishes validation, stale, duplicate, authorization-changed and dependency-unavailable, and offers Refresh record for a stale version.",
     "Users can tell what failed, why, and what to do next, without losing work.",
     "High", "Shared UI primitives; mostly present, formalize"),
    ("UX-014-10", "Register caption and row-identity coverage",
     "Some registers render data grids without an explicit caption or a stable row identity, and the declared total is not uniform.",
     "Extend the shared DataTable contract so every register has a caption, a stable row identity and a declared total, and keep the static contract test green.",
     "Consistent structure for scanning, keyboard use and assistive technology.",
     "Medium", "Shared DataTable"),
    ("UX-014-11", "Responsive register collapse",
     "Dense filter bars and wide tables compress at narrow widths and at high zoom.",
     "Adopt a card-collapse layout for registers below the small breakpoint and let the filter bar stack, verified at 320, 375, 414 and 768 pixels and at 200 percent zoom.",
     "Usable triage on phones and at high zoom.",
     "Medium", "Register surfaces; companion of IMP-013-13"),
    ("UX-014-12", "First-run guidance for an empty system",
     "A new or idle system shows many empty registers with no first-run guidance beyond per-page create actions.",
     "Add a role-aware first-run block on the dashboard that points to the first sensible action and explains that zero counts are real zeros, not failures.",
     "A new user reaches first value without guessing.",
     "Medium", "Dashboard; companion of IMP-013-02"),
    ("UX-014-13", "Navigation memory and in-nav filter",
     "Expanded navigation sections are not remembered beyond the active section, and long sections are hard to scan.",
     "Persist the expanded section between pages and add an in-navigation filter over the ten sections.",
     "Stable orientation across the product.",
     "Low", "Shell navigation; companion of IMP-013-10"),
    ("UX-014-14", "Authenticated accessibility closure",
     "The authenticated keyboard, screen-reader and responsive matrix is not yet proven on the current build; static and unit guards exist.",
     "Complete the authenticated keyboard, screen-reader and responsive matrix at 320, 375, 414, 768, 1024 and 1440, plus 200 percent zoom and text-spacing, on the deployed candidate.",
     "Verifiable accessibility coverage rather than a claim.",
     "High", "Accessibility; companion of IMP-013-15"),
]


def build():
    if not TEMPLATE.exists():
        raise SystemExit(f"missing template: {TEMPLATE}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(str(OUT))
    sect = clear_body(doc)

    def P(t, s="Normal"):
        para(doc, t, s)

    P("Standard Operating Procedure", "Title")
    P("User Experience Standard and Page Interaction Guide", "Heading 1")
    P(
        "Document status: REVIEW DRAFT — not approved | Revision: D0.1 | Prepared by: yazeed (named System Owner) | Prepared: 09 October 2026 | Reviewer: NOT RECORDED | Approver: NOT RECORDED | Effective date: NOT ASSIGNED"
    )
    P(
        "Document reference: SOP-FRS-014. Owner: yazeed (named System Owner). Basis: authenticated live read of the qclevel.top workspace with the existing system-owner account on 09 October 2026, plus the current repository source. This document records the intended user experience of every page and a register of proposed user-experience improvements. Proposed items are NOT IMPLEMENTED and are not approved, effective or verified."
    )

    P("1. Purpose and scope", "Heading 1")
    P(
        "Use User Experience Standard and Page Interaction Guide to understand how a person moves through the QC Operations & Laboratory Management System page by page: how each page opens, what the primary action is, how registers, empty states, loading states, refusals and blockers are presented, and which user-experience improvements are proposed for later application."
    )
    P(
        "This SOP is the user-facing companion to SOP-FRS-013 System Design and Page Map. SOP-FRS-013 states the architecture, navigation model and page inventory; SOP-FRS-014 states the interaction and experience contract that those pages must honour, page by page, and holds the proposed experience improvements."
    )
    P(
        "In scope: the shared page anatomy, the page-by-page experience of the ten navigation sections, the four workspace tools, the entry and audit surfaces, state and empty-state presentation, feedback and refusal behaviour, and the proposed improvements register. Out of scope: server-side business rules, scientific criteria, electronic-signature policy and release governance, which remain owned by SOP-FRS-001 to SOP-FRS-012 and the open owner decisions."
    )
    P(
        "Boundary: this document records experience facts observed on 09 October 2026 and proposed changes. A proposal in Section 7 is not an implemented feature, not an approval, and not a compliance or UAT claim. Nothing here changes permissions, data, controlled vocabulary or the deployed build. A page read proves presentation and navigation only."
    )

    P("2. Responsibilities and prerequisites", "Heading 1")
    P(
        "Named system owner (yazeed): approves this SOP and prioritises the proposals; decides revision, approval and effectivity. Owner approval of this file is separate from in-application registration, signature or effectivity."
    )
    P(
        "Document Control and QMS: review that the described experience matches the deployed interface and the related FRS procedures before controlled adoption."
    )
    P(
        "Domain owners (QC, QMS, Laboratory, Assets, Administration, Governance): confirm page-level wording, blocker wording and the owning decision reference for their pages."
    )
    P(
        "Design and engineering: implement only approved proposals, keep the route, register-surface, mutation-safety and design-governance contract tests green, and update this SOP whenever page behaviour changes."
    )
    P(
        "Access administrator: keep accounts, roles and scopes correct. Visibility of a page is never mutation authority; every action is re-authorized server-side."
    )
    P(
        "Prerequisites: an ACTIVE authenticated account with the required role and scope; the current deployed build; the related SOP-FRS documents; and a browser at a supported viewport. A reviewer who cannot see an action must confirm it is denied server-side, not merely hidden."
    )

    P("3. User experience standard", "Heading 1")
    P("3.1 Shared page anatomy", "Heading 2")
    P(
        "A normal page opens inside the authenticated shell: a skip link, a primary navigation rail grouped into ten sections, a breadcrumb, the current operational scope and the signed-in identity, then the page body, then a separate WORKSPACE TOOLS group. The page body opens with a section kicker, a single level-1 heading and, where relevant, a one-line purpose and a server snapshot time with a Refresh link."
    )
    P(
        "Registers use one shared pattern: a GET filter bar with role=search and a Clear control, applied-filter chips, a table or grid with a caption and a stable row identity, an empty state that separates genuinely empty from filtered-empty, and shared pagination where the register is bounded. Counts and KPIs publish their source, scope and window and distinguish a real zero from an unavailable read."
    )
    P("3.2 Interaction and feedback rules", "Heading 2")
    P(
        "Every page is readable by any ACTIVE authenticated account; only the two owner-only routes /system/health and /system/control-center are refused, server-side with 404. Read visibility never grants an action. An account that can read a register but cannot act sees the state and the true next action without a control it does not hold."
    )
    P(
        "One control per action. A rendered action is the only place an action may be performed, and the server re-checks permission, scope, record state, expected version, separation of duties and any required signature. A missing control means the action is not available to this account, not that the page is broken."
    )
    P(
        "Feedback is classified, not condensed. Validation, stale-version, duplicate-command, authorization-changed and dependency-unavailable outcomes are shown as distinct messages; entered values are preserved on failure; stale intent is never retried against newer data, and a Refresh record control is offered instead. Success returns to the record it wrote."
    )
    P("3.3 The page-by-page experience", "Heading 2")
    P(
        "The following inventory states, for each page, the section, the heading, the purpose line where one is present, the primary action and the interaction controls observed on 09 October 2026. A route that shows no records in this environment is genuinely empty (the database holds no operational records yet); it is not a failed read."
    )
    for route, section, kicker, h1, primary, notes in PAGE_UX:
        P(f"{route} — {section} — kicker \"{kicker}\" — H1 \"{h1}\" — primary: {primary} {notes}")
    P("3.4 State, empty, loading and blocker presentation", "Heading 2")
    P(
        "Structured states such as PASS, FAIL, HOLD, RELEASED and VOID stay literal and are never redefined; human labels are used for everything else. Inspection result and release system state are always shown as separate facts: a scientific PASS is not a release, and HOLD blocks release."
    )
    P(
        "Empty, filtered-empty and unavailable are three different things. A genuinely empty register says so and offers the create action; a filtered-to-empty register says no match and offers to clear filters; an unavailable source says 'Not available' with no number and no link and never collapses to zero. Series and cards use explicit states AVAILABLE, EMPTY, UNAVAILABLE and NOT_SUPPLIED, and only AVAILABLE plots points."
    )
    P(
        "Policy blockers are part of the experience, not an error: where a decision is unresolved, the page states the decision ID and owner, keeps existing records readable, and keeps the blocked action unavailable. The reading is fail-closed by design."
    )

    P("4. Operating procedure", "Heading 1")
    P(
        "4.1 Sign in and read your context. After sign-in, read the signed-in identity and the current operational scope in the shell before acting; every count and register is bounded by that scope."
    )
    P(
        "4.2 Start at the dashboard. Read Needs your attention first, then the work queues, then the operational action counts, then the quarantine flow. Each count opens the exact register it counted, so number and drill-down always agree."
    )
    P(
        "4.3 Use My work today for your personal queue. Read the four groups (Blocked, Overdue, Due today, Assigned to you); groups can overlap and each record appears once, so group counts are not added into a total."
    )
    P(
        "4.4 Open a register. Use the filter bar, read the applied-filter chips, the caption and the row states, and use the shared pagination where present. A register that declares it shows every item in scope is deliberately unbounded."
    )
    P(
        "4.5 Read the record and its state. Open the record, read the facts block, the state, the history and any blocker, then follow the owning domain for the true next action."
    )
    P(
        "4.6 Act only on a rendered action. Perform an action only when its control is present for your account or role and the server accepts it."
    )
    P(
        "4.7 Handle a refusal. Read the classified message, keep your values, and use Refresh record when a stale version is reported. Never re-submit a stale intent against newer data."
    )
    P(
        "4.8 Use the workspace tools. Use Notifications for recipient-scoped updates, Search for authorized records, Account settings to manage your password and session, and Operating guides for daily workflows and blocked-action help."
    )
    P(
        "4.9 Consult the guidance for a state. In Operating guides, match the record and state to the Screen and state guidance table and follow the stated next action; the guide explains that PASS does not release material and that a hidden form is a fail-closed refusal."
    )
    P(
        "4.10 Escalate and hand off by function. Employees escalate blocked work to a Supervisor; Supervisors escalate critical or record decisions to the QCM (Manager); Admin handles access and administration; the named system owner decides platform posture. Prepare a handoff with the route, check time, record reference, current state and the evidence needed."
    )
    P(
        "4.11 Preserve history. Never bypass a state, permission or signature step to move work faster; the audit history records the real transition, and a read is not a completion."
    )

    P("5. Feedback, states and exceptions", "Heading 1")
    P(
        "The dashboard publishes eleven operational action counts, each owned by its domain register and each opening the same filtered read it counted, plus the six ordered quarantine stages, one predicate each."
    )
    P(
        "Exceptions observed in the authenticated read on 09 October 2026: every count was zero because the database holds no operational records yet, not because a source failed; migration schema was DEGRADED with 0045 applied against 0047 shipped and two pending; storage was UNAVAILABLE (not configured in this environment); the outbox was DEGRADED with two pending messages and a worker heartbeat not recorded; and the AI provider was POLICY_DISABLED, so the advisory send control is disabled."
    )
    P(
        "Also observed: the quarantine administration page shows 'Catalogue is temporarily unavailable' while keeping the register readable; the approvals page explains that change-request authorization now requires an account-bound signature while other decision types remain policy-blocked; and the Users page shows one authorized member account. These are honest states, not defects."
    )
    P(
        "PASS and RELEASED remain separate. A receiving item on HOLD cannot be released by any dashboard link, and an approval state is not a release command. A notification or delivery event is not business completion."
    )

    P("6. Records and review checks", "Heading 1")
    P(
        "Record only the evidence needed for the review: route or URL, server snapshot time and timezone, signed-in identity with role and scope, filters applied, record references, observed states, and any error or request reference."
    )
    P(
        "A page read proves presentation and navigation only. It is not proof of approval, release, signature, scientific acceptance or task completion."
    )
    P(
        "Authoring evidence: authenticated live read of thirty-two routes on 09 October 2026 using the existing yazeed session (Administrator, System owner, global scope), with viewport and full-page screenshots attached under SOPs/review-drafts/SOP-FRS-014-attachments/live/, plus a private Stitch proposal project used for visual review candidates."
    )
    P(
        "Before controlled adoption, verify: populated registers and mutations; the proposed items in Section 7 only after implementation; authenticated keyboard, assistive-technology and responsive behaviour; and the migration drift shown on System health. UAT and signature-cycle acceptance remain NOT RUN."
    )
    P(
        "Reviewer: NOT RECORDED. Approver: NOT RECORDED. Reference: SOP-FRS-014. Revision: D0.1. Prepared: 09 October 2026. Status: REVIEW DRAFT — approval, effectivity and in-application registration are NOT RECORDED."
    )

    P("7. Proposed improvements and design changes", "Heading 1")
    P(
        "Every item below is PROPOSED and not implemented. IDs are stable references for tracking and approval. Priority is a design-order suggestion, not an approved schedule. Each item requires owner approval before any code change and must preserve server-side authorization, the PASS/RELEASED separation, the unavailable-is-not-zero rule and the fail-closed refusal of unavailable actions. Items marked as a companion of an IMP-013 item extend the design proposals in SOP-FRS-013 with the user-experience detail."
    )
    for rid, area, obs, prop, effect, prio, owner in IMPROVEMENTS:
        P(f"{rid} — {area} [{prio}]. Observation: {obs} Proposal: {prop} Expected effect: {effect} Owner: {owner}.")
    P(
        "Adoption rule: implement a proposal only through the normal change path, with the owning domain confirming content and the contract tests updated. A proposal is removed from this register only when it is implemented and verified, or explicitly rejected with a recorded reason."
    )

    P("8. References and revision history", "Heading 1")
    P(
        "R1: https://qclevel.top/dashboard; /tasks; /work; /quality/findings; /quality/ncr; /quality/rca; /quality/capa; /reject-reports; /quarantine; /quarantine/receiving; /quarantine/inspections; /quarantine/admin; /laboratory/tests; /assets/equipment; /assets/calibrations; /assets/maintenance; /approvals; /change-requests; /documents; /reports; /ai-advisory; /admin/users; /admin/roles; /admin/permissions; /system/health; /system/control-center; /system/backups; /notifications; /search; /account; /help; /audit. Authenticated live read on 09 October 2026."
    )
    P(
        "R2: SOP-FRS-001 Document Control and Sources; SOP-FRS-002 Work Page; SOP-FRS-003 Quality Pages; SOP-FRS-004 Quarantine Pages; SOP-FRS-005 Laboratory Tests; SOP-FRS-006 Assets Pages; SOP-FRS-007 Governance Pages; SOP-FRS-008 Insights Pages; SOP-FRS-009 Administration Pages; SOP-FRS-010 System Pages; SOP-FRS-011 Workspace Tools; SOP-FRS-012 Electronic Signatures and Approval Authority; SOP-FRS-013 System Design and Page Map."
    )
    P(
        "R3: Documents/UX-WRITING-GUIDE.md; Documents/UI-UX-SPECIFICATION.md; Documents/QC-VISUAL-SYSTEM.md; src/shared/copy/ux-vocabulary.ts; src/shared/copy/format.ts; src/ui/components/** shared UI primitives; tests/unit/ui/register-surface-contract.test.ts; tests/unit/ui/mutation-safety-contract.test.ts; tests/unit/ui/design-governance-contract.test.ts."
    )
    P(
        "R4: Stitch visual proposal project https://stitch.withgoogle.com/projects/2994222140523388486 (private; proposals only, not approved). Local assets: SOPs/review-drafts/SOP-FRS-014-attachments/."
    )
    P(
        "D0.1 | 09 October 2026 | Initial review draft. User-experience standard and page interaction guide for the current build plus fourteen proposed improvements recorded for later application. SOP-FRS-014 assigned by the owner. No approval, effectivity or in-application registration recorded."
    )

    P("Visual operating guide", "Heading 1")
    live_figures = [
        ("01-dashboard.png", "Figure 1. Current dashboard. Attention queue, work queues, action counts and the quarantine flow."),
        ("03-work.png", "Figure 2. Current My work today. Four groups, each with its own empty state."),
        ("10-quarantine-receiving.png", "Figure 3. Current Receiving register. The densest filter bar; inspection result and release state shown separately."),
        ("05-quality-ncr.png", "Figure 4. Current NCR records. A policy-blocker banner names the open decisions."),
        ("13-laboratory-tests.png", "Figure 5. Current Laboratory tests. Filter bar, bounded pagination and a no-records state."),
        ("14-assets-equipment.png", "Figure 6. Current Equipment. Purpose line, state filter and empty state."),
        ("17-approvals.png", "Figure 7. Current My approvals. Signature-policy banner and an empty state that is not evidence of availability."),
        ("22-admin-users.png", "Figure 8. Current Users. One authorized member account with roles and scopes."),
        ("25-system-health.png", "Figure 9. Current System health. Connectivity, readiness and QC acceptance are separated."),
        ("31-help.png", "Figure 10. Current Operating guides. Guidance tables and deep links into the product."),
    ]
    for fname, cap in live_figures:
        figure(doc, LIVE / fname, cap)
    stitch_figures = [
        ("01-register-page.png", "Figure 11. Proposed register-page direction (Stitch). Visual structure only. This is an unimplemented, unreviewed proposal that contains synthetic records and invented navigation/compliance labels (for example a fictional site name and a 21 CFR menu item); none of that content is approved or may enter the product. Only the layout idea is proposed."),
        ("02-record-header.png", "Figure 12. Proposed record-header direction (Stitch). Visual structure only. An unimplemented, unreviewed proposal with synthetic content; it is not an executed or verified workflow and its labels are not approved vocabulary."),
    ]
    for fname, cap in stitch_figures:
        if (STITCH / fname).exists():
            figure(doc, STITCH / fname, cap)
    P(
        "09 October 2026 — Prepared SOP-FRS-014 D0.1 as a review draft. Approval, effective date and in-application registration are NOT RECORDED. Proposed improvements are NOT IMPLEMENTED."
    )

    if sect is not None:
        doc.element.body.append(sect)
    set_header_footer(doc)
    doc.save(str(OUT))
    print("written", OUT)
    print("paragraphs", len(doc.paragraphs), "images", len(doc.inline_shapes))


if __name__ == "__main__":
    build()
