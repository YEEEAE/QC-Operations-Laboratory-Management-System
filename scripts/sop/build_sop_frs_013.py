#!/usr/bin/env python3
"""Author SOP-FRS-013 review draft from the SOP-FRS-011 template.

Preserves the template styles, header/footer, numbering, theme and section
properties; replaces body content; embeds live screenshots and Stitch proposals.
Review-draft only: no approval, effectivity or in-application signature.
"""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/approved/SOP-FRS-011-D0.1-Workspace-Tools.docx"
OUT = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/review-drafts/SOP-FRS-013-D0.1-System-Design-and-Page-Map.docx"
ATTACH = ROOT / "Documents/QC_System_WI_SOP_Pack/SOPs/review-drafts/SOP-FRS-013-attachments"

HEADER = "REVIEW DRAFT — not approved | yazeed (prepared) | 09 October 2026 | Effective date not assigned"
FOOTER = "SOP-FRS-013 | D0.1 | REVIEW DRAFT | Page "


def clear_body(doc: Document):
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        body.remove(sect)
    for child in list(body):
        body.remove(child)
    return sect


def set_header_footer(doc: Document):
    for section in doc.sections:
        for p in section.header.paragraphs:
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = HEADER
            else:
                p.add_run(HEADER)
        for p in section.footer.paragraphs:
            for r in p.runs:
                r.text = ""
            if p.runs:
                p.runs[0].text = FOOTER
            else:
                p.add_run(FOOTER)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if "USING WORKSPACE TOOLS" in p.text:
                            for r in p.runs:
                                r.text = r.text.replace("USING WORKSPACE TOOLS", "SYSTEM DESIGN AND PAGE MAP")
                        if "SOP-FRS-011" in p.text:
                            for r in p.runs:
                                r.text = r.text.replace("SOP-FRS-011", "SOP-FRS-013")


def para(doc, text, style="Normal"):
    return doc.add_paragraph(text, style=style)


def figure(doc, image: Path, caption: str):
    if not image.exists():
        raise FileNotFoundError(image)
    from PIL import Image

    with Image.open(image) as im:
        w_px, h_px = im.size
    max_w, max_h = 16.0, 17.0
    h_at_max_w = max_w * h_px / w_px
    if h_at_max_w <= max_h:
        doc.add_picture(str(image), width=Cm(max_w))
    else:
        doc.add_picture(str(image), height=Cm(max_h))
    para(doc, caption)


def build():
    if not TEMPLATE.exists():
        raise SystemExit(f"missing template: {TEMPLATE}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, OUT)
    doc = Document(str(OUT))
    sect = clear_body(doc)

    def P(t, s="Normal"):
        para(doc, t, s)

    P("Standard Operating Procedure", "Title")
    P("System Design and Page Map", "Heading 1")
    P(
        "Document status: REVIEW DRAFT — not approved | Revision: D0.1 | Prepared by: yazeed (named System Owner) | Prepared: 09 October 2026 | Reviewer: NOT RECORDED | Approver: NOT RECORDED | Effective date: NOT ASSIGNED"
    )
    P(
        "Document reference: SOP-FRS-013. Owner: yazeed (named System Owner). Basis: authenticated live read of the qclevel.top workspace and current repository source on 09 October 2026. This document describes the system design and its pages and records proposed improvements. Proposed items are NOT IMPLEMENTED and are not approved, effective or verified."
    )

    P("1. Purpose and scope", "Heading 1")
    P(
        "Use System Design and Page Map to understand how the QC Operations & Laboratory Management System is organised into domains, pages and shared read surfaces, how each page presents records, states and the true next action, and which design and usability improvements are proposed for later application."
    )
    P(
        "This SOP is the design overview that sits above the page-family SOPs. SOP-FRS-001 records document control; SOP-FRS-002 to SOP-FRS-011 describe the Work, Quality, Quarantine, Laboratory, Assets, Governance, Insights, Administration, System and Workspace Tools pages; SOP-FRS-012 records electronic signature and approval authority. SOP-FRS-013 connects them and states the design contract they share."
    )
    P(
        "In scope: the architecture in one view, the ten-section navigation model, the page inventory and layout conventions, state presentation, and the proposed improvements register. Out of scope: server-side business rules, scientific criteria, signature policy and release governance, which remain owned by the referenced SOPs and the open owner decisions."
    )
    P(
        "Boundary: this document records design facts observed on 09 October 2026 and proposed changes. A proposal in Section 7 is not an implemented feature, not an approval, and not a compliance or UAT claim. Nothing here changes permissions, data or the deployed build."
    )

    P("2. Responsibilities and prerequisites", "Heading 1")
    P(
        "Named system owner (yazeed): approves this SOP and prioritises the proposals; decides revision, approval and effectivity. Owner approval of this file is separate from in-application registration, signature or effectivity."
    )
    P(
        "Document Control and QMS: review that the design description matches the current interface and the related FRS procedures before controlled adoption."
    )
    P(
        "Domain owners (QC, QMS, Laboratory, Assets, Administration): confirm page-level content, policy-blocker wording and the owning decision reference for their pages."
    )
    P(
        "Design and engineering: implement only approved proposals, keep the compile-time route and UI contract tests green, and update this SOP whenever page behaviour changes."
    )
    P(
        "Access administrator: keep accounts, roles and scopes correct. Visibility of a page is never mutation authority; every action is re-authorized server-side."
    )
    P(
        "Prerequisites: an ACTIVE authenticated account; the current deployed build, whose release identity is read from the server artifact; the related SOP-FRS documents; and a browser at a supported viewport."
    )

    P("3. System design and page map", "Heading 1")
    P("3.1 Architecture in one view", "Heading 2")
    P(
        "The application is a server-rendered Astro SSR system with PostgreSQL as the source of truth. Each request is authorized server-side against permission, scope, record state, expected version, separation of duties and, where required, an account-bound electronic signature. Page visibility never grants an action."
    )
    P(
        "Three layers are kept separate: domain modules hold business rules; delivery pages project data safely and contain no SQL or business rules; shared UI primitives give every register the same look and behaviour. The target deployment is Render; a shared shell provides a skip link, the primary navigation, a breadcrumb, the current scope, the signed-in identity, and a separate WORKSPACE TOOLS group."
    )
    P(
        "Two routes are owner-only and refused to every other account: /system/health and /system/control-center. All other normal pages are readable by any ACTIVE authenticated account, while their mutations stay permission-bound."
    )

    P("3.2 Navigation model", "Heading 2")
    P(
        "Primary navigation groups the product into ten sections. Each section expands to its pages; the section containing the current page opens automatically and the current page is marked."
    )
    P(
        "Overview: Dashboard. Work: Tasks, My work today. Quality: Findings, NCR, Root-cause analyses, CAPA, Reject reports. Quarantine: Quarantine overview, Receiving items, Inspection reports, Quarantine administration. Laboratory: Laboratory tests. Assets: Equipment, Calibrations, Maintenance. Governance: My approvals, Change requests, Controlled documents, Releases. Insights: Reports, AI advisory. Administration: Users, Roles, Permissions. System: System health, Control center, Backups."
    )
    P(
        "Separately, WORKSPACE TOOLS groups four cross-domain links: Notifications, Search, Account settings and Operating guides. Audit history is reached from the shell, the Control center and the Reports areas."
    )

    P("3.3 Page layout conventions", "Heading 2")
    P(
        "Every normal page opens with a small section kicker, a level-1 heading, a one-line purpose, and where relevant a server snapshot time with a Refresh link. Registers use a shared pattern: a GET filter bar (role=search) with a Clear control, applied-filter chips, a table with a caption and a stable row identity, and an empty state that separates genuinely empty from filtered-empty."
    )
    P(
        "Counts and KPIs always publish their source, scope and window, and distinguish a real zero from an unavailable read. A refused read shows as 'Not available' with no number and no link; a read that is a real zero is shown as zero with its drill-down link."
    )

    P("3.4 Page inventory (observed 09 October 2026)", "Heading 2")
    P(
        "Overview: /dashboard. Work: /tasks, /work. Quality: /quality/findings, /quality/ncr, /quality/rca, /quality/capa, /reject-reports. Quarantine: /quarantine, /quarantine/receiving, /quarantine/inspections, /quarantine/admin. Laboratory: /laboratory/tests. Assets: /assets/equipment, /assets/calibrations, /assets/maintenance. Governance: /approvals, /change-requests, /documents, and the release-evidence view. Insights: /reports, /reports/[reportCode], /ai-advisory. Administration: /admin/users, /admin/roles, /admin/permissions. System: /system/health, /system/control-center, /system/backups. Workspace tools: /notifications, /search, /account, /help. Audit: /audit. Entry: /login."
    )

    P("3.5 State presentation", "Heading 2")
    P(
        "States are shown through human labels while structured terms such as PASS, FAIL, HOLD, RELEASED and VOID stay literal and are never redefined. Inspection result and release system state are always shown as separate controlled facts: a scientific PASS is not a release, and HOLD blocks release."
    )
    P(
        "A record that is readable with a decision belonging to another role shows its state and the next action without rendering an action the account does not hold; the server refuses that mutation. This fail-closed presentation is by design."
    )

    P("4. Operating procedure", "Heading 1")
    P(
        "4.1 Confirm your context. After sign-in, read the signed-in identity and the current operational scope in the shell before acting. All counts and registers are bounded by that scope."
    )
    P(
        "4.2 Start at the dashboard. Read Needs your attention first, then the work queues, then the action counts. Each count opens the exact register it counted, so the number and its drill-down always agree."
    )
    P(
        "4.3 Move to My work today when you need your personal queue. Read the four groups (Blocked, Overdue, Due today, Assigned to you). Groups can overlap and each record appears once; do not add group counts to form a total."
    )
    P(
        "4.4 Open a register. Use the filter bar, then read the applied-filter chips, the caption and the row states. Use the shared pagination where present; a register that declares it shows every item in scope is genuinely unbounded by design."
    )
    P(
        "4.5 Read the record and its state. Open the record, read the facts block, the state, the history and any blocker, then follow the owning domain route for the true next action."
    )
    P(
        "4.6 Act only on a rendered action. Perform an action only when its control is present for your account or role and the server accepts it. A missing control means the action is not available to you, not that it is broken."
    )
    P(
        "4.7 Handle a refusal. Read the classified message, keep your values, and use Refresh record when a stale version is reported. Never re-submit a stale intent against newer data. Follow the state and troubleshooting guidance in Operating guides."
    )
    P(
        "4.8 Use the workspace tools. Use Notifications for recipient-scoped updates, Search for authorized records, Account settings to manage your password and session, and Operating guides for daily workflows and blocked-action help."
    )
    P(
        "4.9 Consult the state guidance. In Operating guides, match the record and state to the Screen and state guidance table and follow the stated next action; the guide explains that PASS does not release material and that a hidden form is a fail-closed refusal."
    )
    P(
        "4.10 Escalate and hand off by function. Employees escalate blocked work to a Supervisor; Supervisors escalate critical or record decisions to the QCM (Manager); Admin handles access and administration; the named system owner decides platform posture. Prepare a handoff with the route, check time, record reference, current state and the evidence needed."
    )
    P(
        "4.11 Preserve audit history. Never bypass a state, permission or signature step to move work faster. Audit history records the real transition; a read is not a completion."
    )

    P("5. Counts states and exceptions", "Heading 1")
    P(
        "The dashboard publishes eleven action counts, each owned by its domain register and each opening the same filtered read it counted. The quarantine flow publishes six ordered stages, one predicate each, in the order an item moves through quarantine."
    )
    P(
        "Series and cards use explicit states: AVAILABLE, EMPTY, UNAVAILABLE and NOT_SUPPLIED. Only AVAILABLE plots points; an unavailable source never becomes zero."
    )
    P(
        "Exceptions observed in the authenticated read on 09 October 2026: every count was zero because the database holds no operational records yet, not because a source failed; migration schema was DEGRADED with 0045 applied against 0047 shipped and two pending; storage was UNAVAILABLE (not configured in this environment); the outbox was DEGRADED with two pending messages and a worker heartbeat not recorded; and the AI provider was POLICY_DISABLED."
    )
    P(
        "PASS and RELEASED remain separate. A receiving item on HOLD cannot be released by any dashboard link, and an approval state is not a release command."
    )

    P("6. Records and review checks", "Heading 1")
    P(
        "Record only the evidence needed for the review or controlled process: route or URL, server snapshot time and timezone, signed-in identity with role and scope, filters applied, record references, observed states, and any error or request reference."
    )
    P(
        "A page read proves presentation and navigation only. It is not proof of approval, release, signature, scientific acceptance or task completion. Delivery or notification events are not business completion."
    )
    P(
        "Authoring evidence: authenticated live read of more than thirty routes on 09 October 2026 using the existing yazeed session (Administrator, System owner, global scope), plus a private Stitch proposal project used for visual review candidates. Live screenshots and the proposal assets are attached under SOPs/review-drafts/SOP-FRS-013-attachments/."
    )
    P(
        "Before controlled adoption, verify: populated registers and mutations; the proposed items in Section 7 only after implementation; authenticated keyboard, assistive-technology and responsive behaviour; and the migration drift shown on System health. UAT and signature-cycle acceptance remain NOT RUN."
    )
    P(
        "Reviewer: NOT RECORDED. Approver: NOT RECORDED. Reference: SOP-FRS-013. Revision: D0.1. Prepared: 09 October 2026. Status: REVIEW DRAFT — approval, effectivity and in-application registration are NOT RECORDED."
    )

    P("7. Proposed improvements and design changes", "Heading 1")
    P(
        "Every item below is PROPOSED and not implemented. IDs are stable references for tracking and approval. Priority is a design-order suggestion, not an approved schedule. Each item requires owner approval before any code change and must preserve server-side authorization, the PASS/RELEASED separation, the unavailable-is-not-zero rule and the case where a control is refused."
    )
    improvements = [
        ("IMP-013-01", "Dashboard density", "Eleven count cards plus a thirteen-entry Data coverage panel make the first screen heavy.", "Group the counts by domain and move the Data coverage detail into a collapsible disclosure with a one-line summary.", "Lower first-screen cognitive load without losing any source, scope or window.", "High"),
        ("IMP-013-02", "Empty-state economy", "When every count is zero the dashboard shows a long field of zero cards.", "Show a compact single-row strip of zeros with the drill-down links, and expand to full cards when at least one source is non-zero or unavailable.", "A new or idle system reads as calm rather than busy.", "High"),
        ("IMP-013-03", "Filter progressive disclosure", "The receiving register filter bar exposes eight or more controls at once.", "Keep the two or three primary filters visible, move the rest behind More filters, and allow saved filter presets.", "Faster triage on dense registers.", "Medium"),
        ("IMP-013-04", "Register pagination", "Several registers declare that they show every item in scope rather than paging.", "Push the scope predicate into the query and bound the register with the shared pagination, publishing a true total.", "Predictable performance and findability at scale.", "High"),
        ("IMP-013-05", "Policy-blocker component", "Quality registers state policy blockers as ad-hoc prose blocks with different wording.", "Use one consistent blocker card showing the decision ID, the owner and a link, across Findings, NCR, RCA and CAPA.", "Consistent, findable explanation of why an action is unavailable.", "Medium"),
        ("IMP-013-06", "Reject reports zero state", "The reject analytics view renders many empty charts when there is no data.", "Collapse the analytics area into one clear empty state with the create action until data exists.", "Remove the impression of a broken dashboard.", "Medium"),
        ("IMP-013-07", "Approval clarity", "The approvals page explains signature policy in a banner but not per item.", "Show an inline status badge and the blocked reason on each approval item.", "The reviewer sees why a decision is or is not available without guessing.", "Medium"),
        ("IMP-013-08", "Help wayfinding", "Operating guides is a long single page.", "Add a sticky On this page index and in-page search over the guide sections.", "Faster access to the correct guidance.", "Low"),
        ("IMP-013-09", "Command palette", "Ctrl/Cmd+K returns to Search only.", "Extend the palette with recent records and a small set of navigation and quick actions, respecting permissions.", "Fewer clicks for frequent work.", "Medium"),
        ("IMP-013-10", "Navigation memory", "Expanded navigation sections are not remembered between pages beyond the active section.", "Persist the expanded section and add an in-navigation filter for long sections.", "Stable orientation across the ten sections.", "Low"),
        ("IMP-013-11", "Design-token adoption", "Raw colour values remain in a measured set of page positions.", "Replace remaining raw values with shared tokens and keep the compile-time ratchet shrinking.", "Visual consistency and easier theming.", "Medium"),
        ("IMP-013-12", "Restrained motion", "Transitions are minimal today.", "Add short enter and exit transitions for dialogs, drawers and inline feedback only, using transform and opacity, with reduced-motion honoured.", "Perceived quality without distracting from QC work.", "Low"),
        ("IMP-013-13", "Responsive tables", "Dense tables compress at narrow widths.", "Adopt a card-collapse layout for registers below the small breakpoint.", "Usable triage on phones and at high zoom.", "Medium"),
        ("IMP-013-14", "Visual direction", "The Stitch proposal offers a calmer control-room look but contains invented content.", "Adopt only the visual structure and spacing, and map every state label to the approved application vocabulary; discard invented protocols, compliance labels and fictional identities.", "A calmer, more premium shell without inventing facts.", "Medium"),
        ("IMP-013-15", "Accessibility closure", "Authenticated keyboard and assistive-technology coverage is not yet proven.", "Complete the authenticated keyboard, screen-reader and responsive matrix at 320, 375, 414, 768, 1024 and 1440, plus 200 percent zoom and text-spacing.", "Verifiable accessibility coverage for the current build.", "High"),
    ]
    for rid, area, obs, prop, effect, prio in improvements:
        P(f"{rid} — {area} [{prio}]. Observation: {obs} Proposal: {prop} Expected effect: {effect}")
    P(
        "Adoption rule: implement a proposal only through the normal change path, with the owning domain confirming content and the contract tests updated. A proposal is removed from this register only when it is implemented and verified, or explicitly rejected with a recorded reason."
    )

    P("8. References and revision history", "Heading 1")
    P(
        "R1: https://qclevel.top/dashboard; https://qclevel.top/work; https://qclevel.top/quarantine; https://qclevel.top/approvals; https://qclevel.top/help; https://qclevel.top/system/health. Authenticated live read on 09 October 2026."
    )
    P(
        "R2: SOP-FRS-001 Document Control and Sources; SOP-FRS-002 Work Page; SOP-FRS-003 Quality Pages; SOP-FRS-004 Quarantine Pages; SOP-FRS-005 Laboratory Tests; SOP-FRS-006 Assets Pages; SOP-FRS-007 Governance Pages; SOP-FRS-008 Insights Pages; SOP-FRS-009 Administration Pages; SOP-FRS-010 System Pages; SOP-FRS-011 Workspace Tools; SOP-FRS-012 Electronic Signatures and Approval Authority."
    )
    P(
        "R3: src/pages/** page sources; src/ui/components/** shared UI primitives; src/ui/client/dialog.ts; src/shared/copy/ux-vocabulary.ts; src/shared/copy/format.ts; tests/unit/ui/register-surface-contract.test.ts; tests/unit/ui/design-governance-contract.test.ts."
    )
    P(
        "R4: Stitch visual proposal project https://stitch.withgoogle.com/projects/17052786721710239621 (private; proposals only, not approved). Local assets: SOPs/review-drafts/SOP-FRS-013-attachments/."
    )
    P(
        "D0.1 | 09 October 2026 | Initial review draft. Design overview and page map for the current build plus fifteen proposed improvements recorded for later application. SOP-FRS-013 assigned by the owner. No approval, effectivity or in-application registration recorded."
    )

    P("Visual operating guide", "Heading 1")
    figure(doc, ATTACH / "live-01-dashboard.png", "Figure 1. Current dashboard. Counts, work queues and the quarantine flow, each opening the register it counted.")
    figure(doc, ATTACH / "live-02-work.png", "Figure 2. Current My work today. Blocked, overdue, due today and assigned, with the limits of the queue stated.")
    figure(doc, ATTACH / "live-03-quarantine.png", "Figure 3. Current Quarantine overview. Stage distribution with PASS shown separately from release.")
    figure(doc, ATTACH / "live-04-receiving.png", "Figure 4. Current Receiving register. Filter bar, notes and empty state.")
    figure(doc, ATTACH / "01-dashboard.png", "Figure 5. Proposed dashboard direction (Stitch). Visual structure only; contains synthetic content and invented labels that must not enter the product.")
    figure(doc, ATTACH / "02-receiving-register.png", "Figure 6. Proposed receiving register direction (Stitch). Visual structure only; noncanonical state vocabulary is not approved.")
    figure(doc, ATTACH / "03-approval-review.png", "Figure 7. Proposed approval review direction (Stitch). Illustrative only; it is not an executed or verified signature cycle.")
    P(
        "09 October 2026 — Prepared SOP-FRS-013 D0.1 as a review draft. Approval, effective date and in-application registration are NOT RECORDED. Proposed improvements are NOT IMPLEMENTED."
    )

    if sect is not None:
        doc.element.body.append(sect)
    set_header_footer(doc)
    doc.save(str(OUT))
    print("written", OUT)
    print("paragraphs", len(doc.paragraphs), "images", len(doc.inline_shapes))


if __name__ == "__main__":
    build()
