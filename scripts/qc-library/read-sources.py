#!/usr/bin/env python3
"""Bounded, read-only source extraction for QC-DOC-NEON-361.

Archives are inspected in place. No ZIP member is extracted to the filesystem.
The JSONL body output is local working material and must not be treated as a
semantic review, source approval, or controlled-document decision.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import os
import re
import stat
import zipfile
from pathlib import Path, PurePosixPath

MAX_ENTRIES = 20_000
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_TOTAL_BYTES = 1024 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 20_000
MAX_MEMBER_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED = 1024 * 1024 * 1024
MAX_EXPANSION_RATIO = 1_000
MAX_EXTRACTED_CHARS = 2_000_000
MAX_XLSX_SHEETS = 200
MAX_XLSX_CELLS = 2_000_000
MAX_PDF_PAGES = 10_000
CODE_RE = re.compile(r"\b(?:WI(?:-DD)?-[A-Z0-9]+(?:-[A-Z0-9]+)*|F-823-[iT][A-Z0-9-]*|TF-[A-Z0-9-]+|SOP-[A-Z0-9-]+|BS\s*EN\s*ISO\s*\d{3,5}(?:-\d+)*)\b", re.I)
REV_RE = re.compile(r"\b(?:revision|rev\.?|edition|ed\.?|version)\s*[:#.-]?\s*([A-Z0-9]+(?:[.][A-Z0-9]+)*)", re.I)
ISO_RE = re.compile(r"\b(?:BS\s*EN\s*)?ISO\s*\d{3,5}(?:-\d+)*(?::\d{4})?\b", re.I)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def safe_parts(value: str) -> list[str]:
    if not value or "\\" in value or "\0" in value or value.startswith("/"):
        raise ValueError("unsafe manifest path")
    p = PurePosixPath(value)
    if any(part in ("", ".", "..") for part in p.parts):
        raise ValueError("unsafe manifest path")
    return list(p.parts)


def source_path(root: Path, rel: str) -> Path:
    current = root
    for part in safe_parts(rel):
        current = current / part
        info = current.lstat()
        if stat.S_ISLNK(info.st_mode):
            raise ValueError("manifest path traverses symlink")
    resolved = current.resolve(strict=True)
    if resolved == root or root not in resolved.parents:
        raise ValueError("manifest path escapes source root")
    if not resolved.is_file():
        raise ValueError("manifest source is not a regular file")
    return resolved


def clean_text(value: str) -> str:
    value = value.replace("\x00", "")
    return value[:MAX_EXTRACTED_CHARS]


def guard_package(raw: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("document package entry count exceeds bound")
        total = 0
        for info in infos:
            if info.file_size > MAX_MEMBER_BYTES:
                raise ValueError("document package member exceeds bound")
            total += info.file_size
            if total > MAX_ARCHIVE_UNCOMPRESSED:
                raise ValueError("document package uncompressed size exceeds bound")
            if info.file_size and (info.compress_size == 0 or info.file_size / info.compress_size > MAX_EXPANSION_RATIO):
                raise ValueError("document package has a suspicious expansion ratio")


def extract_docx(source: bytes) -> dict:
    from docx import Document
    guard_package(source)
    document = Document(io.BytesIO(source))
    blocks: list[dict] = []
    text: list[str] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        value = paragraph.text.strip()
        if value:
            blocks.append({"location": f"paragraph:{index}", "style": paragraph.style.name, "text": value})
            text.append(value)
    for table_index, table in enumerate(document.tables, 1):
        rows = []
        for row_index, row in enumerate(table.rows, 1):
            values = [cell.text.strip() for cell in row.cells]
            rows.append(values)
            text.append(" | ".join(values))
        blocks.append({"location": f"table:{table_index}", "rows": rows})
    header_footer = []
    for section_index, section in enumerate(document.sections, 1):
        for label, part in (("header", section.header), ("footer", section.footer)):
            values = [p.text.strip() for p in part.paragraphs if p.text.strip()]
            if values:
                header_footer.append({"location": f"section:{section_index}/{label}", "paragraphs": values})
                text.extend(values)
    full = "\n".join(text)
    return {"kind": "DOCX", "text": clean_text(full), "blocks": blocks, "headersFooters": header_footer,
            "contentStatus": "EXTRACTED_COMPLETE" if blocks or header_footer else "EMPTY_DOCUMENT"}


def extract_xlsx(source: bytes) -> dict:
    from openpyxl import load_workbook
    guard_package(source)
    workbook = load_workbook(io.BytesIO(source), read_only=False, data_only=False, keep_links=False)
    sheets = []
    all_text: list[str] = []
    try:
        if len(workbook.worksheets) > MAX_XLSX_SHEETS:
            raise ValueError("workbook sheet count exceeds bound")
        total_cells = 0
        for sheet in workbook.worksheets:
            if sheet.max_row * sheet.max_column > MAX_XLSX_CELLS:
                raise ValueError("worksheet cell grid exceeds bound")
            total_cells += sheet.max_row * sheet.max_column
            if total_cells > MAX_XLSX_CELLS:
                raise ValueError("workbook cell grid exceeds bound")
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    rendered = value.isoformat() if hasattr(value, "isoformat") else str(value)
                    if cell.data_type == "f":
                        rendered = f"FORMULA:{rendered}"
                    cells.append({"cell": cell.coordinate, "value": rendered, "dataType": cell.data_type,
                                  "numberFormat": cell.number_format})
                    all_text.append(f"{sheet.title}!{cell.coordinate}: {rendered}")
            sheets.append({"name": sheet.title, "state": sheet.sheet_state, "dimensions": sheet.calculate_dimension(),
                           "mergedRanges": [str(r) for r in sheet.merged_cells.ranges], "cells": cells})
    finally:
        workbook.close()
    populated = sum(bool(sheet["cells"]) for sheet in sheets)
    status = "EXTRACTED_COMPLETE" if populated else "EMPTY_WORKBOOK"
    return {"kind": "XLSX", "text": clean_text("\n".join(all_text)), "sheets": sheets, "contentStatus": status}


def extract_pdf(source: bytes) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(source), strict=True)
    if reader.is_encrypted:
        raise ValueError("encrypted PDF requires authorized handling")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError("PDF page count exceeds bound")
    pages = []
    parts = []
    for index, page in enumerate(reader.pages, 1):
        value = (page.extract_text() or "").strip()
        if len(value) > MAX_EXTRACTED_CHARS:
            value = value[:MAX_EXTRACTED_CHARS]
        pages.append({"page": index, "text": value})
        parts.append(value)
    full = "\n".join(parts)
    status = "EXTRACTED_COMPLETE" if full.strip() else "NEEDS_SOURCE_RESCAN"
    return {"kind": "PDF", "text": clean_text(full), "pages": pages, "pageCount": len(pages), "contentStatus": status}


def extract_odt(source: bytes) -> dict:
    import xml.etree.ElementTree as ET
    guard_package(source)
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        names = set(archive.namelist())
        if "content.xml" not in names:
            raise ValueError("ODT content.xml is missing")
        xml_bytes = archive.read("content.xml")
    root = ET.fromstring(xml_bytes)
    paragraphs = []
    for node in root.iter():
        if node.tag.endswith("}p") or node.tag.endswith("}h"):
            value = "".join(node.itertext()).strip()
            if value:
                paragraphs.append(value)
    full = "\n".join(paragraphs)
    return {"kind": "ODT", "text": clean_text(full), "paragraphs": paragraphs,
            "contentStatus": "EXTRACTED_COMPLETE" if paragraphs else "EMPTY_DOCUMENT"}


def safe_archive(source: bytes) -> dict:
    members = []
    extracted: list[dict] = []
    total = 0
    with zipfile.ZipFile(io.BytesIO(source)) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ARCHIVE_MEMBERS:
            raise ValueError("archive entry count exceeds bound")
        for info in infos:
            name = info.filename
            path = PurePosixPath(name)
            mode = (info.external_attr >> 16) & 0xFFFF
            is_link = stat.S_ISLNK(mode)
            unsafe = path.is_absolute() or any(part in ("..", ".") for part in path.parts) or "\\" in name or "\0" in name
            total += info.file_size
            if info.file_size > MAX_MEMBER_BYTES or total > MAX_ARCHIVE_UNCOMPRESSED:
                status = "OVERSIZED_NOT_READ"
            elif unsafe:
                status = "UNSAFE_MEMBER_PATH"
            elif is_link:
                status = "SYMLINK_MEMBER_REJECTED"
            elif info.file_size and info.compress_size == 0:
                status = "SUSPICIOUS_COMPRESSION_NOT_READ"
            elif info.compress_size and info.file_size / info.compress_size > MAX_EXPANSION_RATIO:
                status = "SUSPICIOUS_COMPRESSION_NOT_READ"
            elif info.is_dir():
                status = "DIRECTORY_ENTRY"
            else:
                status = "MEMBER_METADATA_ONLY"
                ext = Path(name).suffix.lower()
                if ext in {".docx", ".xlsx", ".pdf", ".odt"}:
                    try:
                        raw = archive.read(info)
                        nested = extract_bytes(raw, ext, allow_archive=False)
                        extracted.append({"archiveMember": name, "sha256": hashlib.sha256(raw).hexdigest(), "extraction": nested})
                        status = "MEMBER_CONTENT_EXTRACTED"
                    except Exception as exc:  # source exceptions are data, not execution instructions
                        status = "MEMBER_EXTRACTION_FAILED"
                        extracted.append({"archiveMember": name, "error": type(exc).__name__})
            members.append({"name": name, "compressedBytes": info.compress_size, "uncompressedBytes": info.file_size,
                            "crc": f"{info.CRC:08x}", "isSymlink": is_link, "status": status})
    return {"kind": "ZIP", "members": members, "extractedMemberContent": extracted,
            "memberCount": len(members), "uncompressedBytes": total, "contentStatus": "ARCHIVE_INSPECTED"}


def extract_bytes(raw: bytes, ext: str, allow_archive: bool = True) -> dict:
    if ext == ".docx": return extract_docx(raw)
    if ext == ".xlsx": return extract_xlsx(raw)
    if ext == ".pdf": return extract_pdf(raw)
    if ext == ".odt": return extract_odt(raw)
    if ext == ".zip" and allow_archive: return safe_archive(raw)
    if ext == ".doc": return {"kind": "LEGACY_WORD", "text": "", "contentStatus": "NEEDS_SOURCE_RESCAN"}
    return {"kind": "UNSUPPORTED", "text": "", "contentStatus": "NEEDS_SOURCE_RESCAN"}


def source_revision(text: str) -> str | None:
    match = REV_RE.search(text[:250_000])
    return match.group(1).upper() if match else None


def title_candidate(text: str, filename: str) -> str:
    for line in text.splitlines():
        candidate = re.sub(r"\s+", " ", line).strip(" |\t:-")
        if len(candidate) >= 4 and not CODE_RE.fullmatch(candidate) and not REV_RE.fullmatch(candidate):
            return candidate[:512]
    return Path(filename).stem


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-root", default="Documents/QC_Library_Source_Import")
    parser.add_argument("--output-dir", default=".ci-results/QC-DOC-NEON-361")
    args = parser.parse_args()
    root = Path(args.source_root).resolve(strict=True)
    out = Path(args.output_dir).resolve()
    repo = Path.cwd().resolve()
    allowed = repo / ".ci-results"
    if out != allowed and allowed not in out.parents:
        raise SystemExit("Output must remain under .ci-results.")
    manifest = json.loads((root / "SOURCE_MANIFEST.json").read_text(encoding="utf-8"))
    records = manifest.get("file_records")
    if not isinstance(records, list) or len(records) > MAX_ENTRIES:
        raise SystemExit("Manifest entry list is absent or exceeds the entry-count bound.")
    out.mkdir(parents=True, exist_ok=True)
    content_path = out / "SOURCE_CONTENT_EXTRACTION.jsonl"
    ledger_path = out / "SOURCE_DIGITIZATION_LEDGER.csv"
    exception_path = out / "SOURCE_EXCEPTION_REGISTER.csv"
    ledger = []
    exceptions = []
    summary = {"manifestRows": len(records), "present": 0, "hashMatched": 0, "extractedComplete": 0,
               "needsRescan": 0, "unsupportedFormats": 0, "extractionErrors": 0, "unavailable": 0, "conflicts": 0,
               "sourceBytesRead": 0, "archivesInspected": 0}
    seen_bytes = 0
    with content_path.open("w", encoding="utf-8") as body_out:
        for row_no, entry in enumerate(records, 1):
            rel = entry.get("export_relative_path", "")
            row = {"manifest_row": row_no, "source_path": rel, "original_filename": entry.get("file_name", ""),
                   "document_code_filename": entry.get("document_code_filename", ""),
                   "document_type": entry.get("classification", ""), "product_item_code": "NOT_CONFIRMED",
                   "source_revision_candidate": entry.get("revision_filename", ""), "header_revision_candidate": "",
                   "filename_revision": entry.get("revision_filename", ""), "document_title_candidate": "",
                   "related_wi_candidates": "", "related_form_candidates": "", "applicable_iso_candidates": "",
                   "workflow_owner": "NOT_ESTABLISHED", "source_page_sheet": "SEE_EXTRACTION_JSONL",
                   "digitization_status": "NOT_REVIEWED", "approval_status": entry.get("content_approval_status", "NOT_APPROVED"),
                   "license_status": "LICENSE_REVIEW_REQUIRED" if entry.get("classification") == "STANDARD" else "NOT_APPLICABLE",
                   "standard_reference_candidate": entry.get("document_code_filename", "") if entry.get("classification") == "STANDARD" else "",
                   "source_sha256": entry.get("sha256", ""), "actual_sha256": "", "actual_size_bytes": "",
                   "integrity_status": "UNAVAILABLE", "content_extraction_status": "NOT_READ"}
            try:
                path = source_path(root, rel)
                size = path.stat().st_size
                summary["present"] += 1
                if size > MAX_FILE_BYTES or seen_bytes + size > MAX_TOTAL_BYTES:
                    raise ValueError("source exceeds configured byte bound")
                seen_bytes += size
                actual = sha256(path)
                row["actual_sha256"] = actual
                row["actual_size_bytes"] = size
                row["integrity_status"] = "MATCH" if actual == str(entry.get("sha256", "")).lower() and size == int(entry.get("file_size_bytes", -1)) else "SOURCE_CONFLICT"
                if row["integrity_status"] != "MATCH":
                    summary["conflicts"] += 1
                    exceptions.append({"manifest_row": row_no, "source_path": rel, "filename": entry.get("file_name", ""),
                                       "classification": entry.get("classification", ""), "status": "SOURCE_CONFLICT",
                                       "expected_sha256": entry.get("sha256", ""), "actual_sha256": actual,
                                       "expected_size_bytes": entry.get("file_size_bytes", ""), "actual_size_bytes": size})
                    body_out.write(json.dumps({"manifestRow": row_no, "sourcePath": rel, "sha256": actual,
                                              "originalFilename": entry.get("file_name", ""),
                                              "extraction": {"kind": "NOT_READ", "text": "",
                                                             "contentStatus": "SOURCE_CONFLICT"}},
                                             ensure_ascii=False) + "\n")
                    ledger.append(row)
                    continue
                summary["hashMatched"] += 1
                if entry.get("classification") == "STANDARD":
                    row["content_extraction_status"] = "LICENSE_REVIEW_REQUIRED"
                    row["digitization_status"] = "LICENSE_REVIEW_REQUIRED"
                    body_out.write(json.dumps({"manifestRow": row_no, "sourcePath": rel, "sha256": actual,
                                              "originalFilename": entry.get("file_name", ""),
                                              "extraction": {"kind": "LICENSE_RESTRICTED", "text": "",
                                                             "contentStatus": "LICENSE_REVIEW_REQUIRED"}},
                                             ensure_ascii=False) + "\n")
                    ledger.append(row)
                    continue
                with path.open("rb") as stream:
                    raw = stream.read(MAX_FILE_BYTES + 1)
                extraction = extract_bytes(raw, path.suffix.lower())
                if extraction.get("kind") == "ZIP": summary["archivesInspected"] += 1
                status = extraction.get("contentStatus", "NEEDS_SOURCE_RESCAN")
                row["content_extraction_status"] = status
                row["digitization_status"] = "SOURCE_CONTENT_EXTRACTED_REQUIRES_REVIEW" if status == "EXTRACTED_COMPLETE" else status
                text = extraction.get("text", "")
                codes = sorted(set(CODE_RE.findall(text)))
                iso_refs = sorted(set(ISO_RE.findall(text)))
                row["header_revision_candidate"] = source_revision(text) or ""
                row["document_title_candidate"] = title_candidate(text, entry.get("file_name", ""))
                row["related_wi_candidates"] = "; ".join(code for code in codes if code.upper().startswith("WI-"))
                row["related_form_candidates"] = "; ".join(code for code in codes if code.upper().startswith("F-823-"))
                row["applicable_iso_candidates"] = "; ".join(iso_refs)
                if status == "EXTRACTED_COMPLETE": summary["extractedComplete"] += 1
                elif status == "NEEDS_SOURCE_RESCAN": summary["needsRescan"] += 1
                elif status != "ARCHIVE_INSPECTED": summary["unsupportedFormats"] += 1
                body_out.write(json.dumps({"manifestRow": row_no, "sourcePath": rel, "sha256": actual,
                                          "originalFilename": entry.get("file_name", ""), "extraction": extraction},
                                         ensure_ascii=False) + "\n")
            except FileNotFoundError:
                summary["unavailable"] += 1
                exceptions.append({"manifest_row": row_no, "source_path": rel, "filename": entry.get("file_name", ""),
                                   "classification": entry.get("classification", ""), "status": "SOURCE_NOT_AVAILABLE",
                                   "expected_sha256": entry.get("sha256", ""), "actual_sha256": "",
                                   "expected_size_bytes": entry.get("file_size_bytes", ""), "actual_size_bytes": ""})
                body_out.write(json.dumps({"manifestRow": row_no, "sourcePath": rel, "sha256": "",
                                          "originalFilename": entry.get("file_name", ""),
                                          "extraction": {"kind": "NOT_AVAILABLE", "text": "",
                                                         "contentStatus": "SOURCE_NOT_AVAILABLE"}},
                                         ensure_ascii=False) + "\n")
            except Exception as exc:
                summary["extractionErrors"] += 1
                row["content_extraction_status"] = "EXTRACTION_ERROR"
                row["digitization_status"] = "NEEDS_SOURCE_RESCAN"
                exceptions.append({"manifest_row": row_no, "source_path": rel, "filename": entry.get("file_name", ""),
                                   "classification": entry.get("classification", ""), "status": "NEEDS_SOURCE_RESCAN",
                                   "expected_sha256": entry.get("sha256", ""), "actual_sha256": row["actual_sha256"],
                                   "expected_size_bytes": entry.get("file_size_bytes", ""), "actual_size_bytes": row["actual_size_bytes"],
                                   "reason_code": type(exc).__name__})
                body_out.write(json.dumps({"manifestRow": row_no, "sourcePath": rel,
                                          "sha256": row["actual_sha256"],
                                          "originalFilename": entry.get("file_name", ""),
                                          "extraction": {"kind": "EXTRACTION_ERROR", "text": "",
                                                         "contentStatus": "NEEDS_SOURCE_RESCAN",
                                                         "errorClass": type(exc).__name__}},
                                         ensure_ascii=False) + "\n")
            ledger.append(row)
    with ledger_path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(ledger[0]) if ledger else [])
        writer.writeheader(); writer.writerows(ledger)
    with exception_path.open("w", newline="", encoding="utf-8-sig") as stream:
        fields = ["manifest_row", "source_path", "filename", "classification", "status", "license_state", "standard_reference_candidate", "expected_sha256", "actual_sha256", "expected_size_bytes", "actual_size_bytes", "reason_code"]
        for exception in exceptions:
            entry = records[exception["manifest_row"] - 1]
            exception["license_state"] = "LICENSE_REVIEW_REQUIRED" if entry.get("classification") == "STANDARD" else "NOT_APPLICABLE"
            exception["standard_reference_candidate"] = entry.get("document_code_filename", "") if entry.get("classification") == "STANDARD" else ""
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader(); writer.writerows(exceptions)
    summary["sourceBytesRead"] = seen_bytes
    (out / "source-reading-summary.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
